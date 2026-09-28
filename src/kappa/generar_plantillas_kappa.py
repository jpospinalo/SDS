#!/usr/bin/env python3
# ADVERTENCIA - script heredado (historico, no reproducible tal cual desde el repositorio).
# Fue escrito para una carpeta de trabajo plana (RAIZ = carpeta del script) que contenia un
# eval.csv y otros insumos propios. Ese eval.csv NO es data/processed/eval.csv: era una copia
# antigua de febrero. Se conserva como registro del metodo. Para reproducir la concordancia
# use:  python3 src/kappa/verificacion_concordancia.py data/external/eval_1.csv \
#         data/external/eval_2.csv data/external/eval_original.csv
"""
Plantillas de codificación para la validación externa del instrumento
=====================================================================
Manuscrito : Documentary Quality Assessment of Safety Data Sheets in the
             Colombian Paint Sector Using Data Analytics
Revista    : ACS Chemical Health & Safety — ID hs-2026-00106u

Genera tres archivos en formato largo (una fila = un juicio), listos para
calcular kappa ponderado por pares y alfa de Krippendorff sin transformaciones
adicionales:

  eval_original.csv  Codificación consensuada original de las 15 FDS.
                     NO se modifica: es el punto de comparación.
  eval_1.csv         Plantilla en blanco para el evaluador externo 1.
  eval_2.csv         Plantilla en blanco para el evaluador externo 2.

Las tres tienen las mismas claves (FDS_ID, ITEM), de modo que se unen
directamente por esas dos columnas.

Cada plantilla viene ordenada según la secuencia asignada a ese evaluador en
orden_codificacion.csv. Los órdenes son distintos a propósito: si ambos
codifican en la misma secuencia, los efectos de aprendizaje y de fatiga se
alinean e inflan el acuerdo de forma espuria.

Los cuatro ítems de la rama de sustancias puras de la Sección 3 (3.1.1 y sus
subítems) no aparecen: los 50 productos del corpus son mezclas y esa rama no
aplica. Quedan 71 ítems evaluables por FDS, es decir 1 065 juicios por
evaluador.

Uso
---
    python3 generar_plantillas_kappa.py

Requiere: pandas, openpyxl, y en el mismo directorio eval.csv, la herramienta
en xlsx, submuestra_kappa_15.csv y orden_codificacion.csv.
"""

import re
from pathlib import Path

import openpyxl
import pandas as pd

RAIZ = Path(__file__).resolve().parent
SALIDA = RAIZ / "kappa"

HERRAMIENTA = "Anexo digital No. 1. Herramienta para verificación de los estándares de comunicación del peligro en las FDS.xlsx"

# La matriz codifica los subítems de la Sección 10 como 10_2_n; el instrumento
# los numera 10_n. Se mapean para poder recuperar la etiqueta de cada uno.
ALIAS = {f"10_2_{i}": f"10_{i}" for i in range(1, 7)}

# Vocabulario normalizado de las plantillas (el archivo original usa
# abreviaturas distintas, que aquí se traducen).
PRESENCIA = {"Presente": "PRESENTE", "No Present": "AUSENTE"}
CALIDAD = {
    "Confiable": "CONFIABLE",
    "Conf_CR": "CONFIABLE_CON_RESTRICCIONES",
    "NO_Conf": "NO_CONFIABLE",
}

VALORES_PRESENCIA = "PRESENTE | AUSENTE"
VALORES_CALIDAD = "CONFIABLE | CONFIABLE_CON_RESTRICCIONES | NO_CONFIABLE"


def etiquetas(ruta_xlsx: Path) -> dict[str, str]:
    """Nombre de cada ítem, leído de las hojas de sección del instrumento."""
    wb = openpyxl.load_workbook(ruta_xlsx, read_only=True, data_only=True)
    nombres: dict[str, str] = {}
    for hoja in wb.sheetnames:
        m = re.match(r"Secci[oó]n\s*(\d+)", hoja.strip())
        if not m:
            continue
        seccion = m.group(1)
        for orden, nombre in wb[hoja].iter_rows(min_row=7, max_col=2, values_only=True):
            if not isinstance(orden, str) or not nombre:
                continue
            orden = orden.strip().rstrip(".")
            if not re.fullmatch(r"\d+(\.\d+)*", orden):
                continue
            codigo = orden.replace(".", "_")
            if codigo.split("_")[0] == seccion:
                nombres.setdefault(codigo, str(nombre).strip())
    return nombres


def formato_largo(df: pd.DataFrame, ids: pd.Series, nombres: dict[str, str]) -> pd.DataFrame:
    """Una fila por juicio: FDS x ítem, con presencia y calidad originales."""
    columnas = [c.split("_", 1)[1] for c in df.columns if c.startswith("Item_")]
    filas = []
    for pos, fila in df.iterrows():
        for cod in columnas:
            calidad = fila.get(f"Calidad_{cod}")
            if pd.isna(calidad):  # rama de sustancias puras: no aplica
                continue
            filas.append(
                {
                    "FDS_ID": ids.iloc[pos],
                    "SECCION": f"S{cod.split('_')[0]}",
                    "ITEM": cod.replace("_", "."),
                    "NOMBRE_ITEM": nombres.get(ALIAS.get(cod, cod), ""),
                    "PRESENCIA": PRESENCIA.get(fila.get(f"Item_{cod}"), ""),
                    "CALIDAD": CALIDAD.get(calidad, ""),
                }
            )
    largo = pd.DataFrame(filas)
    largo["_s"] = largo.SECCION.str[1:].astype(int)
    largo["_i"] = largo.ITEM.map(lambda x: tuple(int(p) for p in x.split(".")))
    return largo.sort_values(["FDS_ID", "_s", "_i"]).drop(columns=["_s", "_i"])


def plantilla(largo: pd.DataFrame, secuencia: list[str], evaluador: int) -> pd.DataFrame:
    """Plantilla en blanco, ordenada según la secuencia de ese evaluador."""
    orden = {fds: i + 1 for i, fds in enumerate(secuencia)}
    t = largo[largo.FDS_ID.isin(secuencia)].copy()
    t["ORDEN_FDS"] = t.FDS_ID.map(orden)
    t["PRESENCIA"] = ""
    t["CALIDAD"] = ""
    t["OBSERVACION"] = ""
    t["EVALUADOR"] = evaluador
    t = t.sort_values(["ORDEN_FDS", "SECCION", "ITEM"], key=lambda s: s if s.name != "SECCION" else s.str[1:].astype(int))
    return t[
        [
            "EVALUADOR",
            "ORDEN_FDS",
            "FDS_ID",
            "SECCION",
            "ITEM",
            "NOMBRE_ITEM",
            "PRESENCIA",
            "CALIDAD",
            "OBSERVACION",
        ]
    ]


def main() -> None:
    SALIDA.mkdir(exist_ok=True)

    df = pd.read_csv(RAIZ / "eval.csv", sep=";", nrows=50, encoding="utf-8-sig")
    sel = pd.read_csv(RAIZ / "submuestra_kappa_15.csv")
    ordenes = pd.read_csv(RAIZ / "orden_codificacion.csv")
    nombres = etiquetas(RAIZ / HERRAMIENTA)

    ids = df["Nombre del Producto"].apply(
        lambda n: "F0{}-P0{}-{}".format(
            re.search(r"FABRICANTE\s*(\d+)", n, re.I).group(1),
            re.search(r"PRODUCTO\s*(\d+)", n, re.I).group(1),
            "I" if "INDUSTRIAL" in n.upper() else "D",
        )
    )

    largo = formato_largo(df, ids, nombres)
    original = largo[largo.FDS_ID.isin(sel.ID)].copy()
    original.insert(0, "EVALUADOR", "ORIGINAL")

    original.to_csv(SALIDA / "eval_original.csv", index=False)
    for i in (1, 2):
        plantilla(largo, list(ordenes[f"Evaluador_{i}"]), i).to_csv(
            SALIDA / f"eval_{i}.csv", index=False
        )

    (SALIDA / "LEEME.txt").write_text(
        "\n".join(
            [
                "VALIDACIÓN EXTERNA DEL INSTRUMENTO — ARCHIVOS DE CODIFICACIÓN",
                "=" * 62,
                "",
                "eval_original.csv   Codificación consensuada original de las 15 FDS.",
                "                    No modificar. Es el punto de comparación.",
                "eval_1.csv          Plantilla del evaluador externo 1.",
                "eval_2.csv          Plantilla del evaluador externo 2.",
                "",
                f"Filas por archivo: {len(original)} ({original.FDS_ID.nunique()} FDS x 71 ítems).",
                "",
                "COLUMNAS QUE DILIGENCIA EL EVALUADOR",
                "-" * 62,
                f"PRESENCIA     {VALORES_PRESENCIA}",
                f"CALIDAD       {VALORES_CALIDAD}",
                "OBSERVACION   Texto libre. Anotar aquí cualquier duda sobre cómo",
                "              aplicar el criterio. Ese registro alimenta el ejemplo",
                "              trabajado y la discusión del kappa por sección.",
                "",
                "No modificar EVALUADOR, ORDEN_FDS, FDS_ID, SECCION, ITEM ni",
                "NOMBRE_ITEM: son las claves con las que se unen los tres archivos.",
                "",
                "INDICACIONES",
                "-" * 62,
                "1. Codificar en el orden indicado por ORDEN_FDS. Cada evaluador",
                "   tiene una secuencia distinta, a propósito.",
                "2. Trabajar de forma independiente. No consultar la codificación",
                "   original ni la del otro evaluador.",
                "3. Usar la herramienta en xlsx (criterios y glosario) como",
                "   referencia para cada ítem.",
                "4. Antes de empezar, calibrar con las dos FDS de fds_calibracion.csv,",
                "   que no forman parte de esta submuestra y se excluyen del cálculo.",
                "",
                "NOTA SOBRE LOS 71 ÍTEMS",
                "-" * 62,
                "El instrumento tiene 75 ítems, pero cuatro corresponden a la rama de",
                "sustancias puras de la Sección 3 (3.1.1 y subítems). Los 50 productos",
                "del corpus son mezclas, así que esa rama no aplica y no aparece en",
                "estos archivos.",
            ]
        )
        + "\n",
        encoding="utf-8",
    )

    print(f"Carpeta: {SALIDA}")
    print(f"  eval_original.csv  {len(original):5d} filas (diligenciado)")
    for i in (1, 2):
        n = len(pd.read_csv(SALIDA / f"eval_{i}.csv"))
        print(f"  eval_{i}.csv         {n:5d} filas (en blanco)")
    print("  LEEME.txt")


if __name__ == "__main__":
    main()
